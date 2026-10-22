"""Create small synthetic document/image fixtures. No personal files are read."""
from pathlib import Path
import argparse
import subprocess
import tempfile

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as SlideInches
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parent
doc = Document()
doc.add_heading("PrintBridge Office fixture", 0)
doc.add_paragraph("中文测试：文档转换、分页与打印预览。")
doc.add_paragraph("Synthetic page one. No physical printing in automated tests.")
doc.add_page_break()
doc.add_heading("Second page", 1)
doc.add_paragraph("The second page must remain a separate page.")
doc.sections[0].page_width = Inches(8.27)
doc.sections[0].page_height = Inches(11.69)
doc.save(root / "sample.docx")

book = Workbook()
sheet = book.active
sheet.title = "数据"
sheet.append(["项目", "数量", "金额"])
sheet.append(["打印测试", 2, 10])
sheet.append(["Synthetic", 3, 12])
sheet.column_dimensions["A"].width = 24
sheet.print_area = "A1:C3"
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
book.save(root / "sample.xlsx")

slides = Presentation()
for title in ["PrintBridge 第一页", "PrintBridge 第二页"]:
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(SlideInches(1), SlideInches(2), SlideInches(8), SlideInches(2))
    box.text_frame.text = "Synthetic presentation / 中文预览测试"
slides.save(root / "sample.pptx")

image = Image.new("RGB", (640, 400), "#ffffff")
draw = ImageDraw.Draw(image)
draw.rectangle((30, 30, 610, 140), fill="#087f73")
draw.text((50, 65), "PrintBridge / TOP", fill="white", font_size=32)
draw.rectangle((30, 180, 300, 360), fill="#d35b45")
draw.text((55, 240), "LEFT", fill="white", font_size=32)
for extension in ["png", "jpg", "webp", "bmp", "gif"]:
    image.save(root / f"sample.{extension}")
second = Image.new("RGB", (400, 640), "#dbe9f1")
ImageDraw.Draw(second).text((40, 120), "TIFF page 2", fill="#202628", font_size=28)
image.save(root / "sample.tiff", save_all=True, append_images=[second])
exif = Image.Exif()
exif[274] = 6
image.save(root / "orientation.jpg", exif=exif)
(root / "sample.txt").write_text("PrintBridge 文本测试\nSynthetic text document.\n", encoding="utf-8")
(root / "sample.csv").write_text("项目,数量\n打印,2\nSynthetic,3\n", encoding="utf-8")
(root / "sample.rtf").write_text(r"{\rtf1\ansi PrintBridge synthetic RTF.\par Second paragraph.}", encoding="ascii")

parser = argparse.ArgumentParser()
parser.add_argument("--office", help="Optional soffice command to generate legacy/ODF fixtures")
args = parser.parse_args()
if args.office:
    for source, target in [("docx", "doc:MS Word 97"), ("xlsx", "xls:MS Excel 97"),
                           ("pptx", "ppt:MS PowerPoint 97"), ("docx", "odt"), ("xlsx", "ods"), ("pptx", "odp")]:
        with tempfile.TemporaryDirectory(prefix="printbridge-fixture-") as profile:
            subprocess.run([args.office, "-env:UserInstallation=" + Path(profile).as_uri(),
                            "--headless", "--convert-to", target, "--outdir", str(root),
                            str(root / ("sample." + source))], check=True, timeout=90,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        assert (root / ("sample." + target.split(":")[0])).is_file()
